"""Gera dados_exemplo/base_de_dados.xlsx com dados 100% FICTÍCIOS (seed fixa).

Uso:  pip install openpyxl  &&  python gerar_dados_exemplo.py
"""
import random
from datetime import date, timedelta
from pathlib import Path

from openpyxl import Workbook

random.seed(42)

MESES = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho", "Julho",
         "Agosto", "Setembro", "Outubro", "Novembro", "Dezembro"]
COLUNAS = ["Núm_Mês", "Mês", "Ano", "Data", "Bandeira", "Cartão", "Natureza",
           "Descrição", "Valor", "Departamento", "Status"]

BANDEIRAS = ["Visa", "Mastercard"]
CARTOES = [1111, 2222, 3333, 4444]  # identificadores fictícios
DEPARTAMENTOS = ["Financeiro", "TI", "Marketing", "Operações", "Eventos"]
# natureza -> (descrições genéricas, faixa de valor)
NATUREZAS = {
    "Transporte":   (["Passagem aérea", "Aplicativo de transporte", "Aluguel de veículo"], (80, 2200)),
    "Hospedagem":   (["Hotel", "Diária de hospedagem"], (250, 3500)),
    "Tecnologia":   (["Assinatura de software", "Licença de ferramenta", "Equipamento"], (50, 4000)),
    "Alimentação":  (["Refeição", "Coffee break", "Almoço de equipe"], (30, 900)),
    "Serviços":     (["Serviço de consultoria", "Taxa de inscrição"], (200, 5000)),
    "Material":     (["Material de escritório", "Material gráfico"], (40, 1500)),
}
STATUS = ["Validado"] * 8 + ["Pendente"] * 2

N = 220
fim = date.today()
inicio = fim - timedelta(days=180)

linhas = []
for _ in range(N):
    d = inicio + timedelta(days=random.randint(0, (fim - inicio).days))
    nat = random.choice(list(NATUREZAS))
    descs, (lo, hi) = NATUREZAS[nat]
    valor = round(random.triangular(lo, hi, lo + (hi - lo) * 0.25), 2)
    linhas.append([d.month, MESES[d.month - 1], d.year, d, random.choice(BANDEIRAS),
                   random.choice(CARTOES), nat, random.choice(descs), valor,
                   random.choice(DEPARTAMENTOS), random.choice(STATUS)])
linhas.sort(key=lambda r: r[3])

wb = Workbook()
ws = wb.active
ws.title = "Base Dados"
ws.append(COLUNAS)
for r in linhas:
    ws.append(r)
for row in ws.iter_rows(min_row=2, min_col=4, max_col=4):
    row[0].number_format = "DD/MM/YYYY"
for row in ws.iter_rows(min_row=2, min_col=9, max_col=9):
    row[0].number_format = "#,##0.00"

saida = Path(__file__).parent / "dados_exemplo"
saida.mkdir(exist_ok=True)
wb.save(saida / "base_de_dados.xlsx")
print(f"{len(linhas)} lançamentos fictícios gravados em {saida / 'base_de_dados.xlsx'}")
